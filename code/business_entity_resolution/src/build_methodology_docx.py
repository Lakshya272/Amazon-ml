"""Convert the editable methodology Markdown to an editable Word document.

Usage: python build_methodology_docx.py Documentation_template.md Documentation_template.docx
Requires python-docx (documentation tooling only).
"""

import re
import sys
from docx import Document
from docx.shared import Inches, Pt


def add_inline(paragraph, value):
    for part in re.split(r"(\*\*.*?\*\*|`[^`]*`)", value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1]).font.name = "Consolas"
        else:
            paragraph.add_run(part)


def main(source, target):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(10)
    for raw in open(source, encoding="utf-8"):
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("- "):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_inline(doc.add_paragraph(), line)
    doc.save(target)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    main(sys.argv[1], sys.argv[2])
